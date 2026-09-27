"""Nexus additions: hybrid semantic retrieval, OCR, the six-agent pipeline with root-cause tracing, analytics, exports and reports."""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
import time

import pytest

import studyhub.web.app as appmod
from studyhub import agents, extract, ocr, retrieval, semantic
from studyhub.db import open_db
from test_studyhub_api import is_error, signed_in  # noqa: F401
from test_studyhub_api_quiz import answer_all, start
from test_studyhub_web import env  # noqa: F401


# ------------------------------------------------------------------------------------------------ helpers

class ToyEmbedder:
    """A deterministic bag-of-concepts embedder: synonyms map to the same axis, so a paraphrase is 'close'."""
    name = "toy-embedder"
    CONCEPTS = [("stack", "lifo", "last", "pile"), ("queue", "fifo", "first", "line"), ("tree", "root", "leaf", "branch"),
                ("recursion", "recursive", "itself", "base"), ("graph", "vertex", "edge", "node")]

    def embed(self, texts):
        out = []
        for t in texts:
            words = re.findall(r"[a-z]+", t.lower())
            v = [sum(1.0 for w in words if w in c) for c in self.CONCEPTS] + [0.01]
            out.append(v)
        return out


@pytest.fixture
def toy(monkeypatch):
    monkeypatch.setenv("STUDYHUB_SEMANTIC", "on")
    semantic.set_embedder(ToyEmbedder())
    yield
    semantic.set_embedder(None)


def seed_graph(sid: int) -> dict[str, int]:
    """Three topics: Recursion <- Trees <- Graphs, 4 questions each (answer index 1)."""
    store = open_db()
    db = store.db
    ids = {}
    doc = db.execute("INSERT INTO documents(subject_id,kind,title,source,sha256,bytes,pages,status,warnings,created_at) VALUES (?,?,?,?,?,?,?,?,?,?)",
                     (sid, "txt", "DS notes", "ds.txt", "0" * 64, 10, None, "parsed", "[]", time.time())).lastrowid
    for i, name in enumerate(["Recursion", "Trees", "Graphs"]):
        ids[name] = db.execute("INSERT INTO topics(subject_id,name,path,ordinal,origin) VALUES (?,?,?,?,'manual')", (sid, name, f"DS > {name}", i)).lastrowid
        for j in range(4):
            db.execute("INSERT INTO mcq_items(subject_id,topic_id,topic_path,question,options,answer_index,explanation,quote,key,created_at,difficulty) "
                       "VALUES (?,?,?,?,?,?,?,?,?,?,?)", (sid, ids[name], f"DS > {name}", f"{name} q{j}?", json.dumps(["a", "b", "c", "d"]), 1, "why", "quote",
                                                       f"{name}{j}", time.time(), ["easy", "medium", "hard", "medium"][j]))
        db.execute("INSERT INTO chunks(subject_id,document_id,topic_id,ordinal,heading_path,text,sha256) VALUES (?,?,?,?,?,?,?)",
                   (sid, doc, ids[name], i, name, f"{name} explained in plain words for the tutor.", hashlib.sha256(name.encode()).hexdigest()))
    db.execute("INSERT INTO topic_prereqs(topic_id,prereq_id,confirmed,origin) VALUES (?,?,1,'manual')", (ids["Trees"], ids["Recursion"]))
    db.execute("INSERT INTO topic_prereqs(topic_id,prereq_id,confirmed,origin) VALUES (?,?,1,'manual')", (ids["Graphs"], ids["Trees"]))
    store.close()
    return ids


# ------------------------------------------------------------------------------------------------ root cause (pure)

def test_root_cause_walks_back_to_the_deepest_weak_foundation():
    parents = {3: [2], 2: [1], 1: []}
    rc = agents.find_root_cause(parents, 3, {1: 0.2, 2: 0.3, 3: 0.1}, 0.7)
    assert rc["root"] == 1 and rc["chain"] == [1, 2, 3] and rc["depth"] == 2


def test_root_cause_stops_at_a_solid_foundation_and_survives_cycles():
    parents = {3: [2], 2: [1, 3], 1: [3]}
    rc = agents.find_root_cause(parents, 3, {1: 0.2, 2: 0.9, 3: 0.1}, 0.7)
    assert rc["root"] == 3 and rc["chain"] == [3]            # 2 is solid, so the walk stops there


# ------------------------------------------------------------------------------------------------ semantic search

def test_hybrid_search_finds_a_paraphrase_that_keyword_search_ranks_low(env, toy):  # noqa: F811
    a = signed_in("alice")
    sid = a.subject("DS")
    text = ("# Stacks\n\nA stack is a LIFO structure: the last item pushed is the first popped, like a pile of plates.\n\n"
            "# Queues\n\nA queue serves items in arrival order, FIFO, like people waiting in a line at a counter.\n")
    assert a.upload(sid, "ds.txt", text.encode()).status_code == 201
    store = open_db()
    try:
        db = store.db
        n = db.execute("SELECT COUNT(*) FROM chunk_embeddings").fetchone()[0]
        if n == 0:                                             # the upload indexed in the background thread: wait for it
            for _ in range(50):
                time.sleep(0.05)
                n = db.execute("SELECT COUNT(*) FROM chunk_embeddings").fetchone()[0]
                if n:
                    break
        assert n >= 2
        uid = db.execute("SELECT id FROM users WHERE username='alice'").fetchone()[0]
        hits = retrieval.search(db, uid, sid, "which structure is last in first out, a pile?", k=3)
        assert hits and "stack" in hits[0].text.lower() and hits[0].similarity is not None
        # scoped: bob's search of his own subject never sees alice's passages
        assert semantic.search(db, uid + 999, sid, "pile") == []
    finally:
        store.close()


def test_semantic_search_is_off_without_an_embedder(env, monkeypatch):  # noqa: F811
    monkeypatch.setenv("STUDYHUB_SEMANTIC", "off")
    assert semantic.embedder() is None and semantic.search(None, 1, 1, "x") == []


def test_vectors_are_stored_normalised_and_compare_by_cosine():
    v = semantic.unpack(semantic.pack([3.0, 4.0]))
    assert math.isclose(v[0], 0.6, abs_tol=1e-6) and math.isclose(v[1], 0.8, abs_tol=1e-6)


# ------------------------------------------------------------------------------------------------ OCR

class FakeReader:
    def readtext(self, arr, detail=1, paragraph=False):
        return [([[0, 0], [100, 0], [100, 20], [0, 20]], "PHOTOSYNTHESIS", 0.95),
                ([[0, 30], [300, 30], [300, 50], [0, 50]], "Plants turn light into chemical energy.", 0.9),
                ([[0, 55], [300, 55], [300, 75], [0, 75]], "Chlorophyll absorbs red and blue light.", 0.88)]


def _png() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (320, 90), "white").save(buf, format="PNG")
    return buf.getvalue()


@pytest.fixture
def fake_ocr(monkeypatch):
    monkeypatch.setenv("STUDYHUB_OCR", "easyocr")
    monkeypatch.setattr(ocr, "_have", lambda m: True)
    monkeypatch.setattr(ocr, "_reader", FakeReader())
    monkeypatch.setattr(ocr, "_engine", "easyocr")
    monkeypatch.setattr(ocr, "_failed", None)


def test_a_photo_of_notes_is_read_with_ocr_into_headed_paragraphs(fake_ocr):
    ex = extract.extract("notes.png", _png())
    assert ex.kind == "image" and ex.status == "parsed" and ex.ocr_pages == 1
    assert ex.blocks[0].level == 1 and ex.blocks[0].text == "PHOTOSYNTHESIS"
    assert "Plants turn light into chemical energy." in ex.blocks[1].text
    assert any("EasyOCR" in w and "confidence" in w for w in ex.warnings)


def test_a_photo_is_refused_plainly_when_no_ocr_engine_is_installed(monkeypatch):
    monkeypatch.setenv("STUDYHUB_OCR", "off")
    with pytest.raises(extract.ExtractError, match="OCR"):
        extract.extract("notes.png", _png())


def test_an_uploaded_photo_becomes_searchable_material_and_its_image_is_served(env, fake_ocr):  # noqa: F811
    a = signed_in("alice")
    sid = a.subject("Biology")
    r = a.upload(sid, "notes.png", _png())
    assert r.status_code == 201, r.text
    doc = r.json()["document"]
    assert doc["kind"] == "image" and doc["ocr_pages"] == 1 and doc["chunks"] >= 1
    f = a.req("GET", f"/subjects/{sid}/materials/{doc['id']}/file")
    assert f.status_code == 200 and f.headers["content-type"] == "image/png"
    s = a.req("GET", f"/subjects/{sid}/search", params={"q": "chlorophyll light"}).json()
    assert s["results"] and "Chlorophyll" in s["results"][0]["text"]


# ------------------------------------------------------------------------------------------------ agents, analytics, exports

def quiz_world(a):
    sid = a.subject("Data Structures")
    ids = seed_graph(sid)
    aid = start(a, sid)
    answer_all(a, sid, aid, chosen=0)                        # everything wrong: every topic is weak
    a.req("POST", f"/subjects/{sid}/quiz/attempts/{aid}/finish")
    return sid, aid, ids


def test_the_agent_pipeline_runs_after_a_quiz_and_traces_misses_to_the_root(env):  # noqa: F811
    a = signed_in("alice")
    sid, aid, ids = quiz_world(a)
    r = a.req("GET", f"/subjects/{sid}/agents", params={"attempt_id": aid})
    assert r.status_code == 200, r.text
    run = r.json()["run"]
    assert [s["agent"] for s in run["steps"]] == ["evaluator", "analytics", "predictor", "planner", "tutor", "mentor"]
    ev = run["steps"][0]
    assert ev["status"] == "done" and ev["details"]["attempt_id"] == aid
    graphs = next(m for m in ev["details"]["misses"] if m["topic"] == "Graphs")
    assert graphs["root"] == "Recursion" and graphs["chain"] == ["Recursion", "Trees", "Graphs"]
    assert ev["details"]["root_causes"][0]["topic"] == "Recursion"
    assert {e["from"] for e in run["graph"]["edges"]} >= {ids["Recursion"], ids["Trees"]}
    # idempotent per quiz; a manual run adds history
    again = a.req("GET", f"/subjects/{sid}/agents", params={"attempt_id": aid}).json()
    assert again["run"]["id"] == run["id"]
    manual = a.req("POST", f"/subjects/{sid}/agents/run").json()
    assert manual["run"]["trigger"] == "manual" and len(manual["history"]) == 2


def test_analytics_overview_tables_and_every_export_format(env):  # noqa: F811
    a = signed_in("alice")
    sid, aid, _ = quiz_world(a)
    ov = a.req("GET", "/analytics", params={"days": 14}).json()
    assert ov["kpis"]["quizzes"] == 1 and ov["kpis"]["answered"] == 12 and ov["kpis"]["accuracy"] == 0
    assert len(ov["timeline"]) == 14 and ov["timeline"][-1]["answered"] == 12
    assert ov["subjects"][0]["name"] == "Data Structures" and ov["subjects"][0]["theta"] < 0
    assert sum(sum(r) for r in ov["habit"]) == 12
    t = a.req("GET", "/analytics/tables/attempts").json()
    assert t["rows"][0]["id"] == aid and t["columns"][0]["key"] == "id"
    topics = a.req("GET", "/analytics/tables/topics", params={"subject_id": sid}).json()["rows"]
    assert {r["topic"] for r in topics} == {"Recursion", "Trees", "Graphs"}
    csv = a.req("GET", "/analytics/tables/attempts.csv")
    assert csv.status_code == 200 and csv.text.lstrip("﻿").startswith("Quiz,Subject") and "attachment" in csv.headers["content-disposition"]
    xlsx = a.req("GET", "/analytics/tables/topics.xlsx")
    assert xlsx.status_code == 200 and xlsx.content[:2] == b"PK"
    pdf = a.req("GET", "/analytics/tables/questions.pdf")
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
    wb = a.req("GET", "/analytics/workbook.xlsx")
    from openpyxl import load_workbook
    names = load_workbook(io.BytesIO(wb.content)).sheetnames
    assert names == ["Subjects", "Daily activity", "Topics", "Attempts", "Questions", "Materials"]
    rep = a.req("GET", "/analytics/report.pdf")
    assert rep.status_code == 200 and rep.content[:5] == b"%PDF-" and len(rep.content) > 3000


def test_subject_and_quiz_pdf_reports(env):  # noqa: F811
    a = signed_in("alice")
    sid, aid, _ = quiz_world(a)
    a.req("GET", f"/subjects/{sid}/agents")
    r = a.req("GET", f"/subjects/{sid}/report.pdf")
    assert r.status_code == 200 and r.content[:5] == b"%PDF-"
    from pypdf import PdfReader
    text = "".join(p.extract_text() for p in PdfReader(io.BytesIO(r.content)).pages)
    assert "Data Structures" in text and "Recursion" in text and "What the agents found" in text
    q = a.req("GET", f"/subjects/{sid}/quiz/attempts/{aid}/report.pdf")
    assert q.status_code == 200 and q.content[:5] == b"%PDF-"
    x = a.req("GET", f"/subjects/{sid}/report.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"


def test_nobody_else_can_read_analytics_reports_or_agents_of_a_subject(env):  # noqa: F811
    a = signed_in("alice")
    sid, aid, _ = quiz_world(a)
    b = signed_in("bob")
    for path, params in [(f"/subjects/{sid}/agents", {}), (f"/subjects/{sid}/report.pdf", {}), (f"/subjects/{sid}/report.xlsx", {}),
                         (f"/subjects/{sid}/quiz/attempts/{aid}/report.pdf", {}), ("/analytics", {"subject_id": sid}),
                         ("/analytics/tables/attempts", {"subject_id": sid}), ("/analytics/report.pdf", {"subject_id": sid})]:
        assert is_error(b.req("GET", path, params=params), 404, "not_found"), path
    assert is_error(b.req("POST", f"/subjects/{sid}/agents/run"), 404, "not_found")
    assert b.req("GET", "/analytics").json()["kpis"]["answered"] == 0       # bob's own analytics show nothing of alice's
    assert b.req("GET", "/analytics/tables/attempts").json()["rows"] == []
    anon = signed_in("carol")
    anon.c.cookies.clear()
    assert is_error(anon.req("GET", "/analytics"), 401, "unauthenticated")


def test_system_status_reports_ocr_and_semantic_capabilities(env, toy):  # noqa: F811
    a = signed_in("alice")
    s = a.req("GET", "/system/status").json()
    assert s["semantic"]["enabled"] is True and s["semantic"]["model"] == "toy-embedder"
    assert "engine" in s["ocr"] and "languages" in s["ocr"]


def test_a_quiz_can_be_sized_and_limited_to_one_difficulty(env):  # noqa: F811
    a = signed_in("alice")
    sid = a.subject("DS")
    seed_graph(sid)
    r = a.req("POST", f"/subjects/{sid}/quiz/attempts", json={"count": 3, "difficulty": "easy"})
    assert r.status_code == 201, r.text
    st = a.req("GET", f"/subjects/{sid}/quiz/attempts/{r.json()['id']}").json()
    assert st["total_questions"] == 3
    assert all(q.endswith("q0?") for q in [st["item"]["question"]])          # q0 is the easy one in every topic
    assert is_error(a.req("POST", f"/subjects/{sid}/quiz/attempts", json={"difficulty": "brutal"}), 400, "invalid")
    assert is_error(a.req("POST", f"/subjects/{sid}/quiz/attempts", json={"count": 50}), 422, "validation")


def test_smart_notes_copy_only_real_sentences_and_refresh_in_place(env):  # noqa: F811
    a = signed_in("alice")
    sid = a.subject("DS")
    text = ("# Stacks\n\nA stack is a last-in first-out collection of elements. The push operation adds an element to the top of the stack. "
            "The pop operation removes the element from the top of the stack. Stacks are used to implement function calls and undo features.\n\n"
            "# Queues\n\nA queue is a first-in first-out collection of elements. The enqueue operation adds an element at the rear of the queue.\n")
    assert a.upload(sid, "ds.txt", text.encode()).status_code == 201
    r = a.req("POST", f"/subjects/{sid}/notes/smart", json={})
    assert r.status_code == 201, r.text
    notes = r.json()["notes"]
    assert {n["title"] for n in notes} == {"Smart notes: Stacks", "Smart notes: Queues"} and all(n["source"] == "smart" for n in notes)
    full = a.req("GET", f"/subjects/{sid}/notes/{notes[0]['id']}").json()
    points = [ln[2:].split("  _(")[0] for ln in full["body"].splitlines() if ln.startswith("- ")]
    assert points and all(p in " ".join(text.split()) for p in points)          # every key point is verbatim from the material
    again = a.req("POST", f"/subjects/{sid}/notes/smart", json={}).json()["notes"]
    assert sorted(n["id"] for n in again) == sorted(n["id"] for n in notes)     # refreshed, not duplicated
    b = signed_in("bob")
    assert is_error(b.req("POST", f"/subjects/{sid}/notes/smart", json={}), 404, "not_found")
